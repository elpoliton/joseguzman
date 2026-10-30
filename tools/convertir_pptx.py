"""Convierte un PowerPoint subido desde el panel /admin en una presentación web de la biblioteca.

Uso:  python3 tools/convertir_pptx.py entrada/<id>.json
El .json trae los datos de la ficha y la ruta del .pptx. Resultado:
  ponencias/<subcategoria>/<id>/index.html   presentación (motor deck.js: puntero, pantalla completa, presentador)
  ponencias/<subcategoria>/<id>/diapos/NN.jpg  cada diapositiva en 1920 px
  ponencias/<subcategoria>/<id>/media/        PDF de respaldo y videos comprimidos
  assets/ponencias.js                         ficha agregada o actualizada
Necesita: LibreOffice (soffice), poppler (pdftoppm), ffmpeg y python-pptx.
"""
import sys, os, re, json, glob, html, shutil, tempfile, subprocess, unicodedata, datetime

from pptx import Presentation
from pptx.oxml.ns import qn

R = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
MESES = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
P14 = '{http://schemas.microsoft.com/office/powerpoint/2010/main}'
RNS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
E = lambda s: html.escape(str(s or ''), quote=True)


def slug(s):
    s = unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')[:60] or 'ponencia'


def sh(*cmd, **kw):
    print('$', ' '.join(cmd), flush=True)
    subprocess.run(cmd, check=True, **kw)


def leer_catalogo():
    src = open(os.path.join(R, 'assets', 'ponencias.js'), encoding='utf-8').read()
    return src[:src.index('[')], json.loads(src[src.index('['):src.rindex(']') + 1])


def guardar_catalogo(cab, P):
    open(os.path.join(R, 'assets', 'ponencias.js'), 'w', encoding='utf-8').write(cab + json.dumps(P, ensure_ascii=False, indent=2) + ';\n')


def videos_de(slide, SW, SH):
    out = []
    for shp in slide.shapes:
        el = shp._element
        vf = el.find('.//' + qn('a:videoFile'))
        if vf is None:
            continue
        m = el.find('.//' + P14 + 'media')
        rid = (m.get(RNS + 'embed') if m is not None else None) or vf.get(RNS + 'link')
        try:
            rel = slide.part.rels[rid]
            if rel.is_external:
                print('  ! video enlazado desde fuera del archivo; se omite:', rel.target_ref); continue
            blob = rel.target_part.blob
        except Exception as e:  # noqa
            print('  ! no se pudo leer un video:', e); continue
        out.append({'blob': blob, 'x': shp.left / SW, 'y': shp.top / SH, 'w': shp.width / SW, 'h': shp.height / SH})
    return out


def main(meta_path):
    meta = json.load(open(meta_path, encoding='utf-8'))
    pptx = os.path.join(R, meta['pptx'])
    sid = slug(meta.get('id') or meta['titulo'])
    sub = slug(meta['subcategoria'])
    ruta = f'{sub}/{sid}/'
    out = os.path.join(R, 'ponencias', sub, sid)
    media, diap = os.path.join(out, 'media'), os.path.join(out, 'diapos')
    shutil.rmtree(out, ignore_errors=True)
    os.makedirs(media); os.makedirs(diap)
    tmp = tempfile.mkdtemp()

    # 1) PDF de respaldo e imágenes de las diapositivas (LibreOffice + poppler)
    sh('soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, pptx, timeout=900)
    pdf = glob.glob(os.path.join(tmp, '*.pdf'))[0]
    pdf_name = f'{sid}.pdf'
    shutil.copy(pdf, os.path.join(media, pdf_name))
    sh('pdftoppm', '-jpeg', '-jpegopt', 'quality=86', '-scale-to-x', '1920', '-scale-to-y', '-1', pdf, os.path.join(tmp, 's'))
    imgs = sorted(glob.glob(os.path.join(tmp, 's-*.jpg')), key=lambda f: int(re.findall(r'(\d+)\.jpg$', f)[0]))
    for k, f in enumerate(imgs):
        shutil.move(f, os.path.join(diap, f'{k + 1:02d}.jpg'))
    N = len(imgs)

    # 2) Textos, notas del orador y videos (python-pptx)
    prs = Presentation(pptx)
    SW, SH = prs.slide_width, prs.slide_height
    info = []
    for s in prs.slides:
        if s._element.get('show') == '0':
            continue  # LibreOffice no exporta las diapositivas ocultas
        textos = [t.text_frame.text.strip() for t in s.shapes if t.has_text_frame and t.text_frame.text.strip()]
        notas = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide and s.notes_slide.notes_text_frame else ''
        info.append({'textos': textos, 'notas': notas, 'videos': videos_de(s, SW, SH)})
    if len(info) != N:
        print(f'  ! {len(info)} diapositivas en el PowerPoint y {N} en el PDF; las notas y videos podrían no coincidir')
        info = (info + [{'textos': [], 'notas': '', 'videos': []}] * N)[:N]

    # posición de la imagen dentro del lienzo 1600 × 900 (object-fit: contain)
    asp = SW / SH
    iw, ih = (1600, 1600 / asp) if asp >= 16 / 9 else (900 * asp, 900)
    ox, oy = (1600 - iw) / 2, (900 - ih) / 2

    nvid, secciones, offline = 0, [], []
    for k in range(N):
        d = info[k]
        titulo_d = d['textos'][0].split('\n')[0][:120] if d['textos'] else ''
        img = f'diapos/{k + 1:02d}.jpg'; offline.append(img)
        partes = [f'<img class="full" src="{img}" alt="Diapositiva {k + 1}{": " + E(titulo_d) if titulo_d else ""}">']
        if d['textos']:
            partes.append('<div class="sr">' + ' · '.join(E(t).replace('\n', ' ') for t in d['textos']) + '</div>')
        for v in d['videos'][:1]:  # el motor maneja un video por diapositiva
            nvid += 1
            raw = os.path.join(tmp, f'v{nvid}.bin'); open(raw, 'wb').write(v['blob'])
            vname, pname = f'video-{nvid}.mp4', f'video-{nvid}.jpg'
            sh('ffmpeg', '-v', 'error', '-y', '-i', raw, '-vf', "scale='min(1280,iw)':-2", '-c:v', 'libx264', '-preset', 'veryfast', '-crf', '26',
               '-pix_fmt', 'yuv420p', '-movflags', '+faststart', '-c:a', 'aac', '-b:a', '128k', os.path.join(media, vname))
            sh('ffmpeg', '-v', 'error', '-y', '-ss', '0.5', '-i', os.path.join(media, vname), '-frames:v', '1', '-q:v', '4', os.path.join(media, pname))
            offline += ['media/' + vname, 'media/' + pname]
            st = f'left:{ox + v["x"] * iw:.0f}px;top:{oy + v["y"] * ih:.0f}px;width:{v["w"] * iw:.0f}px;height:{v["h"] * ih:.0f}px'
            partes.append(f'<div class="vbox" style="{st}"><video src="media/{vname}" poster="media/{pname}" preload="metadata" playsinline></video><div class="hint"><span><i data-icon="FaPlay"></i></span></div></div>')
        if d['notas']:
            partes.append('<aside class="notes">' + ''.join(f'<p>{E(p)}</p>' for p in d['notas'].split('\n') if p.strip()) + '</aside>')
        secciones.append('<section class="slide img">\n  ' + '\n  '.join(partes) + '\n</section>')

    # 3) Página de la presentación
    fecha = meta.get('fecha') or datetime.date.today().isoformat()
    y, m = int(fecha[:4]), int(fecha[5:7])
    ficha = {
        'id': sid, 'titulo': meta['titulo'], 'subtitulo': meta.get('subtitulo', ''), 'categoria': meta['categoria'], 'subcategoria': meta['subcategoria'],
        'fecha': fecha, 'fechaTexto': f'{MESES[m - 1]} {y}', 'lugar': meta.get('lugar', ''), 'resumen': meta.get('resumen', ''),
        'etiquetas': [t.strip() for t in meta.get('etiquetas', []) if t.strip()], 'diapositivas': N, 'videos': nvid,
        'ruta': ruta, 'pdf': ruta + 'media/' + pdf_name, 'portada': ruta + 'diapos/01.jpg', 'oculto': bool(meta.get('oculto', False)), 'origen': 'pptx',
    }
    U = 'https://politon.pe/joseguzman/ponencias/' + ruta
    titulo_pag = f'{ficha["titulo"]} | José Manuel Guzmán Valverde'
    desc = ficha['resumen'] or f'{ficha["titulo"]}: ponencia de José Manuel Guzmán Valverde ({ficha["categoria"]} · {ficha["subcategoria"]}).'
    ld = {'@context': 'https://schema.org', '@type': 'PresentationDigitalDocument', 'name': ficha['titulo'], 'about': ficha['etiquetas'], 'inLanguage': 'es-PE',
          'url': U, 'dateCreated': fecha, 'author': {'@type': 'Person', '@id': 'https://politon.pe/joseguzman/#persona', 'name': 'José Manuel Guzmán Valverde', 'url': 'https://politon.pe/joseguzman/'},
          'encoding': {'@type': 'MediaObject', 'contentUrl': U + 'media/' + pdf_name, 'encodingFormat': 'application/pdf'}}
    base = '../../../'
    off = ','.join([base + 'assets/base.css', base + 'assets/deck.css', base + 'assets/deck.js', base + 'assets/icons.js'] + offline)
    page = f'''<!doctype html>
<html lang="es-PE">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{E(titulo_pag)}</title>
<meta name="description" content="{E(desc)}">
<link rel="canonical" href="{U}">
<meta name="author" content="José Manuel Guzmán Valverde">
<meta property="og:type" content="article">
<meta property="og:locale" content="es_PE">
<meta property="og:url" content="{U}">
<meta property="og:title" content="{E(titulo_pag)}">
<meta property="og:description" content="{E(desc)}">
<meta property="og:image" content="{U}diapos/01.jpg">
<meta name="twitter:card" content="summary_large_image">
<link rel="icon" href="{base}favicon.ico" sizes="48x48">
<link rel="icon" href="{base}favicon.svg" type="image/svg+xml">
<link rel="apple-touch-icon" href="{base}assets/img/apple-touch-icon.png">
<meta name="theme-color" content="#0F1E3D">
<script type="application/ld+json">{json.dumps(ld, ensure_ascii=False)}</script>
<link rel="stylesheet" href="{base}assets/base.css">
<link rel="stylesheet" href="{base}assets/deck.css">
</head>
<body class="deck-page" data-root="{base}" data-offline="{E(off)}">
<h1 class="sr">{E(ficha["titulo"])} — José Manuel Guzmán Valverde</h1>
<div class="viewport"><div class="stage">
{chr(10).join(secciones)}
</div></div>
<div class="blackout"></div>
<div class="progress"></div>
<div class="toast" role="status" aria-live="polite"></div>
<nav class="hud" aria-label="Controles de la presentación">
  <a href="../../" title="Volver a las ponencias"><i data-icon="FaArrowLeft"></i> Ponencias</a>
  <span class="sep"></span>
  <button data-do="prev" title="Anterior (←)"><i data-icon="FaChevronLeft"></i><span class="sr">Anterior</span></button>
  <span class="count">1 / {N}</span>
  <button data-do="next" title="Siguiente (→)"><i data-icon="FaChevronRight"></i><span class="sr">Siguiente</span></button>
  <span class="sep"></span>
  <button data-do="full" title="Pantalla completa (F)"><i data-icon="FaExpand"></i> Pantalla completa</button>
  <button data-do="pres" title="Vista del presentador (S)"><i data-icon="FaDesktop"></i> Presentador</button>
  <a href="media/{pdf_name}" download title="Descargar PDF de respaldo"><i data-icon="FaFilePdf"></i> PDF</a>
  <button data-do="offline" title="Guardar para usar sin internet"><i data-icon="FaWifi"></i> Sin conexión</button>
  <button data-do="help" title="Atajos (?)"><i data-icon="FaKeyboard"></i><span class="sr">Atajos</span></button>
</nav>
<div class="help" role="dialog" aria-label="Atajos de teclado"><div class="box">
  <h3>Atajos y puntero</h3>
  <dl>
    <dt>Siguiente</dt><dd><kbd>→</kbd> <kbd>AvPág</kbd> <kbd>Espacio</kbd> o el botón de avanzar del puntero</dd>
    <dt>Anterior</dt><dd><kbd>←</kbd> <kbd>RePág</kbd> o el botón de retroceder del puntero</dd>
    <dt>Videos</dt><dd>1.er clic: el video se abre en grande · 2.º clic: vuelve a la diapositiva · 3.er clic: avanza. <kbd>K</kbd> pausa o reanuda.</dd>
    <dt>Pantalla completa</dt><dd><kbd>F</kbd> o <kbd>F5</kbd> · salir con <kbd>Esc</kbd></dd>
    <dt>Pantalla negra</dt><dd><kbd>B</kbd> o <kbd>.</kbd></dd>
    <dt>Presentador</dt><dd><kbd>S</kbd> abre una ventana con notas, siguiente diapositiva y cronómetro</dd>
    <dt>Ir a una</dt><dd>número + <kbd>Enter</kbd></dd>
  </dl>
  <p style="margin:16px 0 0;display:flex;justify-content:space-between"><a href="../../probar-puntero.html" style="color:var(--navy);font-weight:700">Probar mi puntero</a><button data-close style="all:unset;cursor:pointer;color:var(--mag);font-weight:700">Cerrar</button></p>
</div></div>
<script src="{base}assets/icons.js"></script>
<script src="{base}assets/deck.js"></script>
</body>
</html>
'''
    open(os.path.join(out, 'index.html'), 'w', encoding='utf-8').write(page)

    # 4) Catálogo: agrega o reemplaza la ficha
    cab, P = leer_catalogo()
    P = [p for p in P if p['id'] != sid] + [ficha]
    guardar_catalogo(cab, P)

    # 5) Limpia la entrada
    for f in (pptx, meta_path):
        if os.path.exists(f):
            os.remove(f)
    shutil.rmtree(tmp, ignore_errors=True)
    print(f'Listo: {ruta} · {N} diapositivas · {nvid} videos')


if __name__ == '__main__':
    for a in sys.argv[1:]:
        main(a)
